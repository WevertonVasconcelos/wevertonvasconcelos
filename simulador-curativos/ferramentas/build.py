# Gera PRECIFICACAO_CURATIVOS_v1.1.xlsx a partir do protótipo v1.0,
# reaproveitando os estilos (fontes, cores, bordas, formatos) das células originais.
import sys
from copy import copy
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill, Font, Alignment

SRC, DST = sys.argv[1], sys.argv[2]
ref = openpyxl.load_workbook(SRC)
RS, RC, RP = ref['Simulador'], ref['Configurações'], ref['Planos confirmados']

wb = openpyxl.Workbook()
wb.remove(wb.active)


def style(dst, src, value=None, fmt=None):
    for a in ('font', 'fill', 'border', 'alignment', 'protection'):
        setattr(dst, a, copy(getattr(src, a)))
    dst.number_format = fmt or src.number_format
    if value is not None:
        dst.value = value
    return dst


C = "'Configurações'!"
MOEDA = RS['B6'].number_format
PCT = RC['B7'].number_format

# ---------------------------------------------------------------- Simulador
ws = wb.create_sheet('Simulador')
ws.sheet_view.showGridLines = False
for col, w in {'A': 44, 'B': 18, 'C': 3, 'D': 38, 'E': 30}.items():
    ws.column_dimensions[col].width = w
ws.column_dimensions['E'].hidden = True  # coluna de controle interno

style(ws['A1'], RS['A1'], 'SIMULADOR DE PLANOS DE CURATIVOS')
style(ws['A2'], RS['A2'], 'VÉRAS VASCONCELOS  ·  MEDICINA E SAÚDE')


def secao(r, texto):
    style(ws.cell(r, 1), RS['A4'], texto)
    style(ws.cell(r, 2), RS['B4'])
    ws.row_dimensions[r].height = 18


secao(4, 'DADOS DO PLANO   ·   preencha os campos amarelos')
entradas = [
    (5, 'Nome do paciente', 'B5', 'Alexandre', 'Obrigatório para salvar.'),
    (6, 'Valor total cobrado', 'B6', 1300, 'Receita bruta do plano.'),
    (7, 'Número de curativos', 'B7', 6, 'Número inteiro maior que zero.'),
    (8, 'Local do atendimento', 'B8', 'Clínica Center', None),
    (9, 'Forma de pagamento', 'B9', 'Pix', 'Dinheiro, Pix, Cartão 1x, 2x ou 3x.'),
    (10, 'Custo de material por curativo', 'B10', 60, 'Deixe em branco se não houver.'),
    (11, 'Custo de deslocamento por curativo', 'B6', None, 'Multiplicado pelo nº de curativos. Em branco = 0.'),
    (12, 'Outras despesas do plano', 'B11', None, 'Valor total: insumo extra, laser etc. (sem deslocamento).'),
    (13, 'Despesas fixas alocadas ao plano', 'B11', None, 'Valor total, descontado uma única vez.'),
]
for r, rot, src, val, dica in entradas:
    style(ws.cell(r, 1), RS['A5'], rot)
    style(ws.cell(r, 2), RS[src], val)
    if dica:
        style(ws.cell(r, 4), RS['D10'], dica)
    ws.row_dimensions[r].height = 18.75

# Controle interno (coluna E oculta): mensagem de erro e taxas encontradas
ok = '$E$2<>""'
E2 = (
    '=IF(AND($B$6="",$B$7=""),"Preencha os campos amarelos para ver o resultado.",'
    'IF(OR(NOT(ISNUMBER($B$6)),N($B$6)<=0),"Valor total cobrado inválido: informe um número maior que zero.",'
    'IF(OR(NOT(ISNUMBER($B$7)),N($B$7)<=0,N($B$7)<>INT(N($B$7))),"Número de curativos inválido: informe um número inteiro maior que zero.",'
    f'IF(ISNA(MATCH($B$8,{C}$A$7:$A$8,0)),"Selecione um local de atendimento válido na lista.",'
    f'IF(ISNA(MATCH($B$9,{C}$A$13:$A$17,0)),"Selecione uma forma de pagamento válida na lista.",'
    'IF(NOT(ISNUMBER($E$3)),"Falta a taxa do local "&$B$8&" em Configurações — cálculo suspenso.",'
    'IF(NOT(ISNUMBER($E$4)),"Falta a taxa de "&$B$9&" em Configurações — cálculo e salvamento suspensos até informá-la.",'
    'IF(NOT(ISNUMBER($E$5)),"Falta a alíquota do Simples Nacional em Configurações — cálculo suspenso.",'
    'IF(OR($E$3<0,$E$3>=1,$E$4<0,$E$4>=1,$E$5<0,$E$5>=1),"Há taxa fora do intervalo de 0% a 99,9% em Configurações — revise.",'
    'IF(AND($B$10<>"",OR(NOT(ISNUMBER($B$10)),N($B$10)<0)),"Custo de material inválido: use valor maior ou igual a zero ou deixe em branco.",'
    'IF(AND($B$11<>"",OR(NOT(ISNUMBER($B$11)),N($B$11)<0)),"Custo de deslocamento inválido: use valor maior ou igual a zero ou deixe em branco.",'
    'IF(AND($B$12<>"",OR(NOT(ISNUMBER($B$12)),N($B$12)<0)),"Outras despesas inválidas: use valor maior ou igual a zero ou deixe em branco.",'
    'IF(AND($B$13<>"",OR(NOT(ISNUMBER($B$13)),N($B$13)<0)),"Despesas fixas inválidas: use valor maior ou igual a zero ou deixe em branco.",'
    '"")))))))))))))'
)


def taxa(nomes, valores, chave):
    idx = f'INDEX({valores},MATCH({chave},{nomes},0))'
    return f'=IFERROR(IF({idx}&""="","",{idx}),"")'


ctrl = Font(name='Arial', size=8, color='FF5A5A5A')
for r, v in [(1, 'CONTROLE INTERNO — não editar'), (2, E2),
             (3, taxa(f'{C}$A$7:$A$8', f'{C}$B$7:$B$8', '$B$8')),
             (4, taxa(f'{C}$A$13:$A$17', f'{C}$B$13:$B$17', '$B$9')),
             (5, f'=IF({C}$B$24&""="","",{C}$B$24)')]:
    ws.cell(r, 5, v).font = ctrl

secao(15, 'DESCONTOS E DESPESAS')
pct = lambda cel: f'SUBSTITUTE(TEXT({cel},"0.0%"),".",",")'
linhas = [
    (16, 'Taxa da clínica', '$B$6*$E$3',
     'IF($E$3=0,$B$8&": sem taxa da clínica",' + pct('$E$3') + '&" do valor cobrado ("&$B$8&")")'),
    (17, 'Taxa da operadora de pagamento', '$B$6*$E$4',
     'IF($E$4=0,$B$9&": sem taxa da operadora",$B$9&": "&' + pct('$E$4') + '&" do valor cobrado")'),
    (18, 'Simples Nacional (estimado)', '$B$6*$E$5', pct('$E$5') + '&" da receita bruta — estimativa a confirmar"'),
    (19, 'Materiais', '$B$7*N($B$10)', '$B$7&" curativos × custo de material"'),
    (20, 'Deslocamento', '$B$7*N($B$11)', '$B$7&" curativos × custo de deslocamento"'),
    (21, 'Outras despesas', 'N($B$12)', None),
    (22, 'Despesas fixas alocadas', 'N($B$13)', '"valor único, não multiplicado"'),
]
for r, rot, f, desc in linhas:
    style(ws.cell(r, 1), RS['A14'], rot)
    style(ws.cell(r, 2), RS['B14'], f'=IF({ok},"",{f})')
    if desc:
        style(ws.cell(r, 4), RS['D14'], f'=IF({ok},"",{desc})')
    ws.row_dimensions[r].height = 16.5
style(ws['A23'], RS['A19'], 'TOTAL DE DESCONTOS E DESPESAS')
style(ws['B23'], RS['B19'], f'=IF({ok},"",SUM($B$16:$B$22))')
ws.row_dimensions[23].height = 21.75

secao(25, 'RESULTADO DO PLANO')
for r, rot, f, src in [(26, 'VALOR FINAL  (líquido para a empresa)', '$B$6-$B$23', 'B22'),
                       (27, 'PREÇO LÍQUIDO POR CURATIVO', '$B$26/$B$7', 'B23'),
                       (28, 'MARGEM APÓS AS DESPESAS INFORMADAS', '$B$26/$B$6', 'B24')]:
    style(ws.cell(r, 1), RS['A22'], rot)
    style(ws.cell(r, 2), RS[src], f'=IF({ok},"",{f})')
    ws.row_dimensions[r].height = 25.5
style(ws['D28'], RS['D10'], 'Margem = valor final ÷ valor cobrado (não é markup).')

# Mensagem de situação (A30:D31)
style(ws['A30'], RS['A26'],
      '=IF($E$2<>"",$E$2,IF(TRIM($B$5&"")="","Plano calculado. Informe o nome do paciente para poder salvar.",'
      'IF($B$26<0,"ATENÇÃO: com estes valores o plano fecha com PREJUÍZO. Revise antes de fechar.",'
      '"Plano calculado. Confira e clique em “Salvar simulação” para registrar no histórico.")))')
for c in 'BCD':
    style(ws[f'{c}30'], RS['A26'])
    ws[f'{c}30'].value = None
ws.merge_cells('A30:D31')

style(ws['A33'], RS['A29'],
      'Para registrar: clique em “Salvar simulação” (ao lado do resultado) ou no menu Simulador → Salvar simulação. '
      'O registro vai para a aba “Histórico de simulações” com valores fixos, que não mudam depois. O simulador continua preenchido.')
ws.merge_cells('A33:D34')
ws.row_dimensions[33].height = 15
style(ws['A36'], RS['A32'], 'Taxas e memória de cálculo: aba Configurações.')

# Validações de dados (mensagens em português)
def dv(tipo, sqref, msg, **kw):
    v = DataValidation(type=tipo, allow_blank=True, showErrorMessage=True, errorStyle='stop',
                       errorTitle='Valor inválido', error=msg, **kw)
    v.add(sqref)
    ws.add_data_validation(v)

dv('decimal', 'B6', 'Informe um valor numérico maior que zero.', operator='greaterThan', formula1='0')
dv('whole', 'B7', 'Informe um número inteiro maior que zero.', operator='greaterThanOrEqual', formula1='1')
dv('decimal', 'B10:B13', 'Informe um valor maior ou igual a zero (ou deixe em branco).',
   operator='greaterThanOrEqual', formula1='0')
dv('list', 'B8', 'Escolha um local da lista.', formula1=f"{C}$A$7:$A$8")
dv('list', 'B9', 'Escolha uma forma de pagamento da lista.', formula1=f"{C}$A$13:$A$17")

# Destaque de prejuízo e de mensagens de bloqueio
alerta_fill = PatternFill('solid', start_color='FFFCE4E4', end_color='FFFCE4E4')
ws.conditional_formatting.add('A26:B28', FormulaRule(
    formula=['AND($E$2="",N($B$26)<0)'], fill=alerta_fill, font=Font(bold=True, color='FFC00000')))
ws.conditional_formatting.add('A30:D31', FormulaRule(
    formula=['OR(AND($E$2<>"",NOT(AND($B$6="",$B$7=""))),AND($E$2="",N($B$26)<0))'],
    fill=alerta_fill, font=Font(bold=True, color='FF8C3A3A')))

# ------------------------------------------------------------ Configurações
wc = wb.create_sheet('Configurações')
wc.sheet_view.showGridLines = False
for col, w in {'A': 46, 'B': 16, 'C': 3}.items():
    wc.column_dimensions[col].width = w
style(wc['A1'], RC['A1'], 'CONFIGURAÇÕES — TAXAS E CÁLCULOS')
style(wc['A2'], RC['A2'], 'Véras Vasconcelos Medicina e Saúde  ·  altere apenas as taxas abaixo')


def csec(r, texto, cab):
    style(wc.cell(r, 1), RC['A5'], texto)
    style(wc.cell(r + 1, 1), RC['A6'], cab[0])
    style(wc.cell(r + 1, 2), RC['B6'], cab[1])


def nota(r, texto, src='A9', ate=None):
    style(wc.cell(r, 1), RC[src], texto)
    if ate:
        wc.merge_cells(f'A{r}:B{ate}')
        wc.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top')


csec(5, '1. TAXA DO LOCAL DE ATENDIMENTO', ('Local', 'Taxa sobre o valor cobrado'))
for r, nome, v in [(7, 'Clínica Center', 0.2), (8, 'Domicílio', 0.0)]:
    style(wc.cell(r, 1), RC['A7'], nome)
    style(wc.cell(r, 2), RC['B7'], v)
nota(9, 'A Clínica Center recebe 20% do valor bruto. No domicílio não há essa taxa.')

csec(11, '2. TAXA DA OPERADORA (FORMA DE PAGAMENTO)', ('Forma de pagamento', 'Taxa sobre o valor cobrado'))
for r, nome, v, src in [(13, 'Dinheiro', 0.0, 'B13'), (14, 'Pix', 0.0, 'B13'), (15, 'Cartão 1x', None, 'B14'),
                        (16, 'Cartão 2x', 0.06, 'B15'), (17, 'Cartão 3x', 0.084, 'B16')]:
    style(wc.cell(r, 1), RC['A13'], nome)
    style(wc.cell(r, 2), RC[src], v)
nota(18, 'Cartão 1x está em branco de propósito: a taxa ainda não foi informada. Enquanto estiver vazia, '
         'o Simulador não calcula nem salva planos em 1x (avisa na tela).', ate=19)
nota(20, 'Dinheiro e Pix: sem taxa da operadora · Cartão 2x = 6% (aproximado, informado pelos sócios) · Cartão 3x = 8,4%.',
     ate=21)

csec(23, '3. IMPOSTO', ('Tributo', 'Alíquota'))
style(wc['A24'], RC['A23'], 'Simples Nacional (estimativa)')
style(wc['B24'], RC['B23'], 0.06)
nota(25, 'REGRA DO MVP 1.1: Simples Nacional = receita bruta do plano × alíquota acima, aplicada sobre o valor cobrado '
         'integral (antes das taxas da clínica e da operadora e das despesas), em todas as formas de pagamento, '
         'inclusive dinheiro e Pix. A alíquota de 6% é estimativa e varia por faixa de faturamento — confirme com a '
         'contabilidade e ajuste aqui quando necessário.', src='A24', ate=28)

style(wc['A30'], RC['A29'], '4. COMO O SIMULADOR CALCULA')
regras = [
    'Taxa da clínica = valor cobrado × taxa do local (item 1)',
    'Taxa da operadora = valor cobrado × taxa da forma de pagamento (item 2)',
    'Simples Nacional = valor cobrado (receita bruta) × alíquota (item 3)',
    'Materiais = número de curativos × custo de material por curativo',
    'Deslocamento = número de curativos × custo de deslocamento por curativo',
    'Outras despesas e despesas fixas alocadas = valores totais informados (uma vez)',
    'Despesas totais = clínica + operadora + Simples + materiais + deslocamento + outras + fixas',
    'Valor final = valor cobrado − despesas totais (pode ser negativo = prejuízo)',
    'Preço líquido por curativo = valor final ÷ número de curativos',
    'Margem = valor final ÷ valor cobrado (margem sobre a receita, não markup)',
]
for i, t in enumerate(regras):
    style(wc.cell(31 + i, 1), RC['A30'], '•  ' + t)
nota(42, 'O Simulador deixa os resultados em branco quando falta um campo, quando há valor inválido ou quando a taxa '
         'da forma de pagamento escolhida não está cadastrada aqui — assim nenhum número parece válido sem estar. '
         'O botão “Salvar simulação” faz as mesmas verificações antes de gravar.', ate=44)
style(wc['A46'], RC['A43'], 'Versão 1.1 — 24/10/2026 · salvamento automático no histórico '
                            '(base: v1.0 de 21/09/2026, revisão Dr. Weverton Vasconcelos)')
for r in (18, 20, 25, 42):
    wc.row_dimensions[r].height = 15

# --------------------------------------------------------------- Histórico
wh = wb.create_sheet('Histórico de simulações')
wh.sheet_view.showGridLines = False
style(wh['A1'], RP['A1'], 'HISTÓRICO DE SIMULAÇÕES')
style(wh['A2'], RP['A2'],
      'Registros gravados automaticamente pelo botão “Salvar simulação” (ou menu Simulador → Salvar simulação), '
      'como valores fixos: não mudam quando o Simulador ou as taxas forem alterados. '
      'Não insira linhas nem edite acima da linha 5.')
wh.merge_cells('A2:L3')
wh['A2'].alignment = Alignment(wrap_text=True, vertical='top')
cab = [('ID do registro', 22), ('Data e hora', 18), ('Nome do paciente', 26), ('Receita bruta', 14),
       ('Nº de curativos', 11), ('Local de atendimento', 16), ('Forma de pagamento', 14),
       ('Material por curativo', 13), ('Materiais (total)', 13), ('Deslocamento por curativo', 14),
       ('Deslocamento (total)', 14), ('Outras despesas', 13), ('Despesas fixas alocadas', 14),
       ('Taxa da clínica (%)', 11), ('Taxa da clínica (R$)', 13), ('Taxa da operadora (%)', 11),
       ('Taxa da operadora (R$)', 13), ('Simples Nacional (%)', 11), ('Simples Nacional (R$)', 13),
       ('Total de descontos e despesas', 15), ('Resultado líquido estimado', 15),
       ('Resultado líquido por curativo', 15), ('Margem estimada', 11)]
for i, (t, w) in enumerate(cab, start=1):
    c = style(wh.cell(4, i), RP['A4'], t)
    c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    wh.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
wh.row_dimensions[2].height = 15
wh.row_dimensions[3].height = 15
wh.row_dimensions[4].height = 42
wh.freeze_panes = 'D5'
wh.auto_filter.ref = 'A4:W1000'

wb.active = 0
wb.save(DST)
print('ok', DST)
