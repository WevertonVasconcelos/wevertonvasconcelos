# Recalcula a planilha no LibreOffice para cada cenário e confere os resultados das fórmulas
import openpyxl, subprocess, sys, os, shutil
SRC = sys.argv[1]
CEN = {
 'A_clinica_pix': dict(B5='Paciente A',B6=1000,B7=5,B8='Clínica Center',B9='Pix',B10=50,B11=None,B12=None,B13=None),
 'B_domicilio_2x': dict(B5='Paciente B',B6=1000,B7=5,B8='Domicílio',B9='Cartão 2x',B10=50,B11=20,B12=None,B13=50),
 'C_dinheiro': dict(B5='Paciente C',B6=1000,B7=5,B8='Domicílio',B9='Dinheiro',B10=None,B11=None,B12=None,B13=None),
 'E1_zero_curativos': dict(B5='Paciente E',B6=1000,B7=0,B8='Clínica Center',B9='Pix',B10=50),
 'E2_cartao1x': dict(B5='Paciente E',B6=1000,B7=5,B8='Clínica Center',B9='Cartão 1x',B10=50),
 'F_prejuizo': dict(B5='Paciente F',B6=300,B7=5,B8='Clínica Center',B9='Cartão 3x',B10=50,B11=20,B12=10,B13=15),
 'G_fracionado': dict(B5='X',B6=1000,B7=2.5,B8='Clínica Center',B9='Pix'),
 'H_negativo': dict(B5='X',B6=1000,B7=5,B8='Clínica Center',B9='Pix',B12=-10),
 'I_vazio': dict(B5=None,B6=None,B7=None,B8=None,B9=None,B10=None),
 'J_sem_nome': dict(B5=None,B6=1000,B7=5,B8='Domicílio',B9='Pix'),
}
cells = ['B16','B17','B18','B19','B20','B21','B22','B23','B26','B27','B28','A30','D16','D17','D18','E2']
for nome, vals in CEN.items():
    wb = openpyxl.load_workbook(SRC); ws = wb['Simulador']
    for k,v in vals.items(): ws[k].value = v
    p = f'lo/{nome}.xlsx'; wb.save(p)
    subprocess.run(['soffice','--headless','--calc','--convert-to','xlsx','--outdir','lo/out',p],capture_output=True,check=True)
    r = openpyxl.load_workbook(f'lo/out/{nome}.xlsx', data_only=True)['Simulador']
    print('\n##', nome)
    for c in cells:
        v = r[c].value
        print(f'  {c}: {round(v,6) if isinstance(v,float) else v!r}')
